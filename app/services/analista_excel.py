"""
EXCEL DO ANALISTA — a planilha do período, gerada a partir do banco.

Abas: Resumo (indicadores + script da academia + funil + dia a dia + gráficos),
Análises (IA), Script - dito × ideal, Notas por critério, Técnicas, Leads,
Conversas e Legenda. As abas da IA só aparecem se existir alguma análise no período.

Estilo: cabeçalho roxo em negrito, bordas finas, linhas zebradas, notas e
situações coloridas, filtro e cabeçalho congelado, pronta para imprimir.
"""
import io
from datetime import datetime, timedelta, timezone

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from app.core.config import config
from app.repositories import mensagens_repo, sessoes_repo
from app.services import analista
from app.services.analista_ia import TODOS_CRITERIOS, _temperatura, carregar_metodo
from app.services.analista_metricas import formatar_duracao
from app.services.analista_relatorios import (
    DIAS, FUNIL, MOTIVOS_PT, TERMOMETRO_PLURAL, aderencia_da_conversa, agregar, faixa_aderencia,
)

TZ = config.FUSO_HORARIO

# ---------------------------------------------------------------------------
# Estilo
# ---------------------------------------------------------------------------
ROXO = "4F46E5"
FONTE = "Calibri"
_lado = Side(style="thin", color="BFBFBF")
BORDA = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)
PREENCHE_CABECALHO = PatternFill("solid", start_color=ROXO, end_color=ROXO)
PREENCHE_ZEBRA = PatternFill("solid", start_color="F5F5FA", end_color="F5F5FA")
PREENCHE_SECAO = PatternFill("solid", start_color="EEF2FF", end_color="EEF2FF")

VERDE, AMARELO, VERMELHO = "C6EFCE", "FFEB9C", "FFC7CE"
AZUL, CINZA = "DDEBF7", "EDEDED"
SITUACOES = {                       # texto -> cor
    "Sucesso": VERDE,
    "Em andamento": AZUL,
    "Pede atenção": AMARELO,
    "Esperando resposta": VERMELHO,
    "Desistiu": VERMELHO,
    "Não é venda": CINZA,
}
VERDE_CLARO = "EAF6EC"               # fundo da coluna "Pelo script poderia ser"

# Termômetro do lead: rótulo, o que quer dizer e a cor (do frio, azul, ao quente, laranja)
TERMOMETRO_INFO = {
    "tranquilo": ("Tranquilo", "Só curiosidade, sem urgência; responde pouco e não fala em começar.", "DDEBF7"),
    "incomodado": ("Incomodado", "Sente que precisa mudar (peso, dor, saúde, autoestima), mas ainda não busca ativamente.",
                   "E2EFDA"),
    "pesquisador": ("Pesquisador", "Está buscando informação (modalidades, horários, como funciona).", "FFF2CC"),
    "comparador": ("Comparador", "Compara preço, plano ou condições (\"tá caro\", \"tem mais barato?\").", "FCE4D6"),
    "comprador": ("Comprador", "Quer começar logo: pergunta como matricular, quando pode ir, pede aula ou visita.",
                  "F8CBAD"),
}
# Etapas do script (até onde a conversa chegou)
ETAPAS_INFO = {
    "acolhimento": ("Acolhimento", "A equipe recebeu o cliente, mas o objetivo dele ainda não ficou claro."),
    "descoberta": ("Descoberta", "A equipe ficou sabendo do objetivo, da dor ou do momento do cliente."),
    "foco": ("Próximo passo", "A equipe propôs uma ligação, visita ou aula, mas nada ficou agendado ainda."),
    "agendamento": ("Agendamento", "Ficou agendada ligação, visita, aula experimental ou matrícula, com dia e hora."),
    "nenhuma": ("Nenhuma", "A equipe não chegou a atender de verdade (sem resposta ou resposta solta)."),
}
FAIXAS_INFO = {          # faixa -> (rótulo, intervalo, cor, o que quer dizer)
    "seguiu": ("Seguiu o script", "70% ou mais", VERDE,
               "A equipe fez a maior parte dos passos que tinham chance de acontecer."),
    "em_parte": ("Em parte", "de 40% a 69%", AMARELO,
                 "Fez parte dos passos; o que faltou aparece na aba \"Script - dito × ideal\"."),
    "normal": ("Atendimento normal", "abaixo de 40%", VERMELHO,
               "Atendeu, mas sem seguir o roteiro do script na maioria dos passos."),
}
# (o Excel não aceita ":" no nome de aba, por isso o traço)
ABA_SCRIPT = "Script - dito × ideal"
ROTULOS_FUNIL = {"acolhidos": "Acolhidos", "descobriram": "Descobriram o objetivo",
                 "proximo_passo": "Com próximo passo", "agendaram": "Agendaram"}


def _fill(cor):
    return PatternFill("solid", start_color=cor, end_color=cor)


def _cor_nota(nota):
    return VERDE if nota >= 8 else AMARELO if nota >= 6 else VERMELHO


def _limpo(valor):
    """Texto seguro para a planilha: sem caracteres proibidos no XML e sem virar fórmula."""
    if isinstance(valor, str):
        valor = ILLEGAL_CHARACTERS_RE.sub("", valor)
        if valor.startswith("="):
            valor = " " + valor
    return valor


def _naive(dt):
    return dt.astimezone(TZ).replace(tzinfo=None) if dt else None


def _config_impressao(ws, titulo_linhas="1:1"):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if titulo_linhas:
        ws.print_title_rows = titulo_linhas


def _tabela(ws, cabecalhos, linhas, larguras, formatos=None, centralizar=(), linha_inicial=1):
    """
    Escreve cabeçalho + linhas com o estilo padrão e devolve a última linha usada.
    formatos: {indice_da_coluna(1..): number_format}; centralizar: índices (1..).
    """
    formatos = formatos or {}
    for c, titulo in enumerate(cabecalhos, start=1):
        cel = ws.cell(row=linha_inicial, column=c, value=titulo)
        cel.fill, cel.border = PREENCHE_CABECALHO, BORDA
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF")
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = larguras[c - 1]
    ws.row_dimensions[linha_inicial].height = 32
    r = linha_inicial
    for i, linha in enumerate(linhas):
        r = linha_inicial + 1 + i
        for c, valor in enumerate(linha, start=1):
            cel = ws.cell(row=r, column=c, value=_limpo(valor))
            cel.border = BORDA
            cel.font = Font(name=FONTE, size=10)
            cel.alignment = Alignment(vertical="top", wrap_text=True,
                                      horizontal="center" if c in centralizar else "left")
            if c in formatos:
                cel.number_format = formatos[c]
            if i % 2 == 1:
                cel.fill = PREENCHE_ZEBRA
    ws.freeze_panes = ws.cell(row=linha_inicial + 1, column=3 if len(cabecalhos) > 6 else 1)
    if linhas:
        ws.auto_filter.ref = f"A{linha_inicial}:{get_column_letter(len(cabecalhos))}{r}"
    _config_impressao(ws, f"{linha_inicial}:{linha_inicial}")
    return r


def _situacao(a):
    """Rótulo colorido da conversa na planilha."""
    if a["desfecho"] == "sucesso":
        return "Sucesso"
    if a["desfecho"] == "perdido":
        return "Desistiu"
    if a["desfecho"] == "nao_lead" or not a["eh_lead"]:
        return "Não é venda"
    if a.get("aguardando_desde"):
        return "Esperando resposta"
    return "Pede atenção" if a["precisa_atencao"] else "Em andamento"


def _dia_label(iso):
    d = datetime.fromisoformat(iso)
    return f"{DIAS[d.weekday()][:3]} {d.strftime('%d/%m')}"


# ---------------------------------------------------------------------------
# Abas
# ---------------------------------------------------------------------------
def _aba_resumo(wb, agg, inicio, fim):
    ws = wb.active
    ws.title = "Resumo"
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{config.CLIENT_NAME} — Análise do atendimento"
    ws["A1"].font = Font(name=FONTE, size=18, bold=True, color=ROXO)
    ws["A2"] = f"Período: {_naive(inicio):%d/%m/%Y %H:%M} até {_naive(fim):%d/%m/%Y %H:%M}"
    ws["A2"].font = Font(name=FONTE, size=11, italic=True, color="595959")

    limite = config.LIMITE_SEM_RESPOSTA_MIN
    pct = (f" ({round(100 * agg['respondidas_no_limite'] / agg['com_primeira_resposta'])}%)"
           if agg["com_primeira_resposta"] else "")
    blocos = [
        ("Movimento", [
            ("Conversas com clientes", agg["conversas"]),
            ("Clientes novos", agg["novas"]),
            ("Clientes que voltaram", agg["retomadas"]),
            ("Mensagens dos clientes", agg["msgs_cliente"]),
            ("Mensagens da equipe", agg["msgs_equipe"]),
            ("Clientes antigos chamados de volta pela equipe", agg["followups_da_equipe"]),
        ]),
        ("Resultados (lidos pela IA)", [
            ("Conversas analisadas pela IA", f"{agg['analisadas']} de {agg['analisaveis']}"),
            ("Com sucesso (agendaram ou fecharam)", agg["sucesso"]),
            ("Em andamento", agg["em_andamento"]),
            ("Desistiram", agg["perdido"]),
            ("Não eram venda", agg["nao_lead"]),
            ("Pedem atenção", agg["atencao_total"]),
        ]),
        ("Velocidade da equipe", [
            ("1ª resposta em média (horário útil)", formatar_duracao(agg["primeira_resposta_media_util_seg"])),
            (f"1ª resposta dentro de {limite} min",
             f"{agg['respondidas_no_limite']} de {agg['com_primeira_resposta']}{pct}"),
            ("Maior espera (horário útil)", formatar_duracao(agg["pior_espera_util_seg"])),
            ("Esperando resposta agora", len(agg["sem_resposta"])),
            ("Sem nenhuma resposta da equipe", agg["respondidas_por"].get("ninguem", 0)),
        ]),
        ("Qualidade do atendimento (lida pela IA)", [
            ("Nota média (0 a 10)", agg["nota_media"] if agg["nota_media"] is not None else "-"),
            ("Intuito de venda (0 a 10)", agg["nota_venda"] if agg["nota_venda"] is not None else "-"),
            ("Propôs ligação para o cliente", f"{agg['propos_ligacao']} de {agg['leads']}"),
        ]),
    ]
    sc = agg.get("script") or {}
    if sc.get("conversas") or sc.get("termometro"):
        blocos.append((f"{sc['titulo']} (lido pela IA)", _itens_script(sc)))
    r = 4
    for titulo, itens in blocos:
        r = _secao(ws, r, titulo)
        for rotulo, valor in itens:
            a, b = ws.cell(row=r, column=1, value=rotulo), ws.cell(row=r, column=2, value=_limpo(valor))
            a.font, b.font = Font(name=FONTE, size=10), Font(name=FONTE, size=10, bold=True)
            a.border = b.border = BORDA
            b.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            r += 1
        r += 1
    funil = sc.get("funil") or {}
    if funil.get("com_etapa"):
        r = _secao(ws, r, "Funil do script (até onde as conversas chegaram)")
        for c, titulo in ((1, "Etapa"), (2, "Conversas")):
            cel = ws.cell(row=r, column=c, value=titulo)
            cel.fill, cel.border = PREENCHE_CABECALHO, BORDA
            cel.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
            cel.alignment = Alignment(horizontal="left" if c == 1 else "right")
        r += 1
        base = funil["acolhidos"]
        for k, (chave, _, _, _) in enumerate(FUNIL):
            qtd = funil[chave]
            valor = qtd if k == 0 or not base else f"{qtd} · {round(100 * qtd / base)}%"
            a, b = ws.cell(row=r, column=1, value=ROTULOS_FUNIL[chave]), ws.cell(row=r, column=2, value=valor)
            a.font, b.font = Font(name=FONTE, size=10), Font(name=FONTE, size=10, bold=True)
            a.border = b.border = BORDA
            b.alignment = Alignment(horizontal="right")
            if k % 2 == 1:
                a.fill = b.fill = PREENCHE_ZEBRA
            r += 1
        nota = ws.cell(row=r, column=1, value="% sobre os acolhidos. Quem agendou também passou pelas etapas anteriores.")
        nota.font = Font(name=FONTE, size=9, italic=True, color="595959")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 3

    # dia a dia (colunas D em diante) + gráficos
    dias = agg["por_dia"]
    d0, d1 = _naive(inicio).date(), _naive(fim).date()
    todos = [(d0 + timedelta(days=i)) for i in range((d1 - d0).days + 1)]
    linhas = []
    for d in todos:
        v = dias.get(d.isoformat())
        if v is None and d.weekday() >= 5:
            continue                                  # fim de semana fechado e sem conversa
        v = v or {"conversas": 0, "novas": 0, "retomadas": 0, "sucesso": 0, "em_andamento": 0, "perdido": 0,
                  "nota_media": None, "primeira_resposta_media_util_seg": None}
        prim = v["primeira_resposta_media_util_seg"]
        linhas.append([_dia_label(d.isoformat()), v["conversas"], v["novas"], v["retomadas"], v["sucesso"],
                       v["em_andamento"], v["perdido"], v["nota_media"],
                       round(prim / 60, 1) if prim is not None else None])
    cab = ["Dia", "Conversas", "Novos", "Voltaram", "Sucesso", "Em andamento", "Desistiram", "Nota média",
           "1ª resposta (min)"]
    larg = [12, 11, 9, 10, 9, 13, 11, 10, 14]
    ws.cell(row=4, column=4, value="Dia a dia").font = Font(name=FONTE, size=12, bold=True, underline="single", color=ROXO)
    for c in range(4, 4 + len(cab)):
        ws.cell(row=4, column=c).fill = PREENCHE_SECAO
    fim_tab = _tabela_deslocada(ws, cab, linhas, larg, coluna_inicial=4, linha_inicial=5)
    ws.row_dimensions[5].height = 20          # a linha 5 também é indicador: não deixar o Excel esticar
    if linhas:
        n = len(linhas)
        barras = BarChart()
        barras.type, barras.title = "col", "Conversas por dia"
        barras.add_data(Reference(ws, min_col=5, min_row=5, max_row=5 + n), titles_from_data=True)
        barras.set_categories(Reference(ws, min_col=4, min_row=6, max_row=5 + n))
        barras.height, barras.width, barras.legend = 6.5, 16, None
        barras.varyColors = False
        barras.series[0].graphicalProperties.solidFill = ROXO
        barras.x_axis.delete = barras.y_axis.delete = False
        ws.add_chart(barras, f"D{fim_tab + 2}")
        if any(l[7] is not None for l in linhas):
            linha = LineChart()
            linha.title = "Nota média por dia"
            linha.add_data(Reference(ws, min_col=11, min_row=5, max_row=5 + n), titles_from_data=True)
            linha.set_categories(Reference(ws, min_col=4, min_row=6, max_row=5 + n))
            linha.height, linha.width, linha.legend = 6.5, 16, None
            linha.y_axis.scaling.min, linha.y_axis.scaling.max = 0, 10
            serie = linha.series[0]
            serie.smooth = False
            serie.graphicalProperties.line.solidFill = ROXO
            serie.graphicalProperties.line.width = 28000
            serie.marker.symbol = "circle"
            serie.marker.size = 7
            serie.marker.graphicalProperties.solidFill = ROXO
            linha.x_axis.delete = linha.y_axis.delete = False
            ws.add_chart(linha, f"D{fim_tab + 16}")
    _config_impressao(ws, None)
    ws.page_setup.fitToHeight = 1          # o Resumo cabe numa folha só (com o bloco do script e o funil)


def _secao(ws, r, titulo):
    """Título de seção (roxo, sublinhado, fundo lilás) nas colunas A:B. Devolve a próxima linha."""
    for c in (1, 2):
        ws.cell(row=r, column=c).fill = PREENCHE_SECAO
    ws.cell(row=r, column=1, value=titulo).font = Font(name=FONTE, size=12, bold=True, underline="single", color=ROXO)
    return r + 1


def _itens_script(sc):
    """Linhas (rótulo, valor) do bloco "Script da academia" da aba Resumo."""
    itens = [
        ("Conversas de venda avaliadas no script", sc["conversas"]),
        ("Aderência ao script (passos feitos ÷ oportunidades)",
         f"{sc['aderencia']}%" if sc.get("aderencia") is not None else "-"),
        ("Seguiram o script (70% ou mais)", sc["seguiram"]),
        ("Em parte (40% a 69%)", sc["em_parte"]),
        ("Atendimento normal (abaixo de 40%)", sc["normal"]),
    ]
    itens += [(p["rotulo"], f"{p['feitas']} de {p['oportunidades']}") for p in sc.get("passos") or []]
    if sc.get("mais_seguido"):
        itens.append(("Passo mais seguido", sc["mais_seguido"]["nome"]))
    if sc.get("mais_esquecido"):
        itens.append(("Passo mais esquecido", sc["mais_esquecido"]["nome"]))
    itens += [(f"Termômetro: {TERMOMETRO_PLURAL[t['perfil']]}", t["qtd"]) for t in sc.get("termometro") or []]
    return itens


def _tabela_deslocada(ws, cabecalhos, linhas, larguras, coluna_inicial, linha_inicial):
    """Tabela simples com borda a partir de (linha_inicial, coluna_inicial). Devolve a última linha."""
    for i, titulo in enumerate(cabecalhos):
        c = coluna_inicial + i
        cel = ws.cell(row=linha_inicial, column=c, value=titulo)
        cel.fill, cel.border = PREENCHE_CABECALHO, BORDA
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = larguras[i]
    r = linha_inicial
    for k, linha in enumerate(linhas):
        r = linha_inicial + 1 + k
        for i, valor in enumerate(linha):
            cel = ws.cell(row=r, column=coluna_inicial + i, value=valor)
            cel.border = BORDA
            cel.font = Font(name=FONTE, size=10, bold=(i == 0))
            cel.alignment = Alignment(horizontal="left" if i == 0 else "center")
            if i == 7 and valor is not None:
                cel.number_format = "0.0"
                cel.fill = _fill(_cor_nota(valor))
            elif k % 2 == 1:
                cel.fill = PREENCHE_ZEBRA
    return r


def _aderencia(a, metodo):
    """(proporção 0-1 ou None, faixa) da conversa, para a planilha."""
    if not a["eh_lead"]:
        return None, None
    feitas, faltaram = aderencia_da_conversa(a["dados"], metodo["tecnicas"])
    oport = len(feitas) + len(faltaram)
    return (len(feitas) / oport if oport else None), faixa_aderencia(len(feitas), oport)


def _aba_analises(wb, analises, metodo):
    ws = wb.create_sheet("Análises (IA)")
    cab = ["Data", "Nome", "Telefone", "Gênero", "Cliente", "Interesse", "Situação", "Nota", "Intuito de venda",
           "Intuito de ligação", "1ª resposta (min)", "Termômetro", "Etapa do script", "Aderência ao script (%)",
           "Resumo", "Ponto forte", "Pode melhorar", "Próxima ação", "Mensagem sugerida"]
    linhas, situacoes, notas, temperaturas, faixas = [], [], [], [], []
    for a in analises:
        prim = a["primeira_resposta_util_seg"]
        sit = _situacao(a)
        if sit == "Desistiu" and a["motivo_perda"]:
            sit_txt = f"Desistiu ({MOTIVOS_PT.get(a['motivo_perda'], a['motivo_perda'])})"
        else:
            sit_txt = sit
        situacoes.append(sit)
        notas.append(a["nota_final"])
        temp = _temperatura(a.get("temperatura")) if a["eh_lead"] else None
        temperaturas.append(temp)
        etapa = (a["dados"] or {}).get("etapa_script") if a["eh_lead"] else None
        aderencia, faixa = _aderencia(a, metodo)
        faixas.append(faixa)
        linhas.append([
            _naive(a["inicio"]), a["nome"] or "Sem nome", a["telefone"], a["genero"] or "Indefinido",
            "Novo" if a["nova"] else "Voltou", a["interesse"], sit_txt,
            float(a["nota_final"]) if a["nota_final"] is not None else None,
            float(a["nota_intuito_venda"]) if a["nota_intuito_venda"] is not None else None,
            float(a["nota_ligacao"]) if a["nota_ligacao"] is not None else None,
            round(prim / 60, 1) if prim is not None else None,
            TERMOMETRO_INFO[temp][0] if temp else None,
            ETAPAS_INFO[etapa][0] if etapa in ETAPAS_INFO else None,
            aderencia,
            a["resumo"], a["ponto_forte"], a["pode_melhorar"], a["proxima_acao"], a["mensagem_sugerida"],
        ])
    ultima = _tabela(ws, cab, linhas,
                     [16, 20, 16, 11, 9, 24, 20, 8, 10, 10, 11, 13, 14, 12, 42, 42, 42, 36, 46],
                     formatos={1: "dd/mm/yyyy hh:mm", 8: "0.0", 9: "0.0", 10: "0.0", 14: "0%"},
                     centralizar=(4, 5, 7, 8, 9, 10, 11, 12, 13, 14))
    for i, (sit, nota, temp, faixa) in enumerate(zip(situacoes, notas, temperaturas, faixas)):
        r = 2 + i
        ws.cell(row=r, column=7).fill = _fill(SITUACOES[sit])
        for col in (8, 9, 10):
            v = ws.cell(row=r, column=col).value
            if v is not None:
                ws.cell(row=r, column=col).fill = _fill(_cor_nota(v))
        ws.cell(row=r, column=8).font = Font(name=FONTE, size=10, bold=True)
        if temp:
            ws.cell(row=r, column=12).fill = _fill(TERMOMETRO_INFO[temp][2])
        if faixa:
            ws.cell(row=r, column=14).fill = _fill(FAIXAS_INFO[faixa][2])
            ws.cell(row=r, column=14).font = Font(name=FONTE, size=10, bold=True)
    return ultima


def _aba_script(wb, analises, metodo):
    """Uma linha por momento: o que o cliente disse, o que a equipe respondeu e como seria pelo script."""
    ws = wb.create_sheet(ABA_SCRIPT)
    linhas = []
    for a in analises:
        if not a["eh_lead"]:
            continue
        for m in (a["dados"] or {}).get("momentos") or []:
            if not isinstance(m, dict):
                continue
            tecnica = metodo["tecnicas"].get(m.get("tecnica"))
            nome_tecnica = f"{m.get('tecnica')} — {tecnica['nome']}" if tecnica else m.get("tecnica")
            linhas.append([_naive(a["inicio"]), a["nome"] or "Sem nome", a["telefone"], nome_tecnica,
                           m.get("cliente_disse") or "-", m.get("equipe_respondeu") or "(ficou sem resposta)",
                           m.get("pelo_script")])
    cab = ["Data", "Nome", "Telefone", "Técnica", "O cliente disse", "A equipe respondeu", "Pelo script poderia ser"]
    ultima = _tabela(ws, cab, linhas, [16, 20, 16, 26, 42, 42, 56], formatos={1: "dd/mm/yyyy hh:mm"})
    ws.freeze_panes = "C2"
    for r in range(2, ultima + 1 if linhas else 2):
        cel = ws.cell(row=r, column=7)
        cel.fill = _fill(VERDE_CLARO)
        ws.cell(row=r, column=4).font = Font(name=FONTE, size=10, bold=True)
    if not linhas:
        ws.cell(row=3, column=2, value="Nenhum momento registrado pela IA neste período.").font = Font(
            name=FONTE, italic=True, color="595959")
    return ultima


def _aba_notas(wb, analises):
    ws = wb.create_sheet("Notas por critério")
    chaves = list(TODOS_CRITERIOS)
    cab = ["Data", "Nome", "Telefone", "Nota final"] + [TODOS_CRITERIOS[k] for k in chaves]
    linhas = []
    for a in analises:
        if not a["eh_lead"]:
            continue
        notas = (a["dados"] or {}).get("notas") or {}
        linhas.append([_naive(a["inicio"]), a["nome"] or "Sem nome", a["telefone"],
                       float(a["nota_final"]) if a["nota_final"] is not None else None]
                      + [notas.get(k) if notas.get(k) is not None else "n/a" for k in chaves])
    larguras = [16, 20, 16, 10] + [13] * len(chaves)
    ultima = _tabela(ws, cab, linhas, larguras, formatos={1: "dd/mm/yyyy hh:mm", 4: "0.0"},
                     centralizar=tuple(range(4, 5 + len(chaves))))
    ws.row_dimensions[1].height = 48
    for r in range(2, ultima + 1):
        for c in range(4, 5 + len(chaves)):
            cel = ws.cell(row=r, column=c)
            if isinstance(cel.value, (int, float)):
                cel.fill = _fill(_cor_nota(cel.value))
                cel.number_format = "0.0"
            elif cel.value == "n/a":
                cel.font = Font(name=FONTE, size=9, color="808080", italic=True)
    if linhas:                                           # linha de médias
        r = ultima + 1
        ws.cell(row=r, column=1, value="Média").font = Font(name=FONTE, bold=True)
        for c in range(1, 5 + len(chaves)):
            ws.cell(row=r, column=c).border = Border(top=Side(style="medium", color=ROXO), bottom=_lado,
                                                     left=_lado, right=_lado)
            if c >= 4:
                letra = get_column_letter(c)
                cel = ws.cell(row=r, column=c, value=f'=IFERROR(AVERAGE({letra}2:{letra}{ultima}),"")')
                cel.font, cel.number_format = Font(name=FONTE, bold=True), "0.0"
                cel.alignment = Alignment(horizontal="center")


def _aba_tecnicas(wb, analises, metodo):
    ws = wb.create_sheet("Técnicas")
    aplicadas, reforcar = {}, {}
    leads = [a for a in analises if a["eh_lead"]]
    for a in leads:
        feitas, faltaram = aderencia_da_conversa(a["dados"])       # cada conversa conta 1 vez por técnica
        for t in feitas:
            aplicadas[t] = aplicadas.get(t, 0) + 1
        for t in faltaram:
            reforcar[t] = reforcar.get(t, 0) + 1
    linhas = []
    for tid, t in metodo["tecnicas"].items():
        feitas, faltou = aplicadas.get(tid, 0), reforcar.get(tid, 0)
        linhas.append([tid, t["nome"], t["como"], feitas, faltou,
                       feitas / (feitas + faltou) if feitas + faltou else None])
    _tabela(ws, ["ID", f"Técnica ({metodo['titulo']})", "Como reconhecer", "Aplicada em (conversas)",
                 "A reforçar em (conversas)", "% seguida"], linhas, [7, 34, 70, 14, 14, 11],
            formatos={6: "0%"}, centralizar=(1, 4, 5, 6))
    for r, linha in enumerate(linhas, start=2):
        ws.cell(row=r, column=4).fill = _fill(VERDE) if linha[3] else PatternFill()
        ws.cell(row=r, column=5).fill = _fill(AMARELO) if linha[4] else PatternFill()
        faixa = faixa_aderencia(linha[3], linha[3] + linha[4])
        if faixa:
            ws.cell(row=r, column=6).fill = _fill(FAIXAS_INFO[faixa][2])
            ws.cell(row=r, column=6).font = Font(name=FONTE, size=10, bold=True)
        else:
            ws.cell(row=r, column=6).value = "n/a"
            ws.cell(row=r, column=6).font = Font(name=FONTE, size=9, color="808080", italic=True)
    fonte = ("script da academia (PDF \"Scripts para academias - CRM\")" if metodo["fonte"] == "curso"
             else "boas práticas gerais (reserva)")
    ws.cell(row=len(linhas) + 3, column=2,
            value=f"Base: {len(leads)} conversas de venda analisadas. Fonte das técnicas: {fonte}. "
                  f"% seguida = aplicada ÷ (aplicada + a reforçar); n/a = não teve oportunidade."
            ).font = Font(name=FONTE, italic=True, color="595959")


def _aba_leads(wb, sessoes):
    ws = wb.create_sheet("Leads")
    linhas = []
    for s in sessoes:
        linhas.append([
            s["telefone"], s["nome"] or "Sem nome", _naive(s["inicio"]), _naive(s["fim"]),
            analista._min(s["duracao_seg"]), analista._sim_nao(s["nova"]),
            "Cliente" if s["iniciada_por"] == "cliente" else "Equipe",
            s["msgs_cliente"], s["msgs_atendente"], s["msgs_bot"],
            analista._min(s["primeira_resposta_util_seg"]), analista._min(s["primeira_resposta_seg"]),
            analista._min(s["tempo_medio_resposta_util_seg"]), analista._min(s["tempo_medio_resposta_seg"]),
            analista._min(s["tempo_max_resposta_util_seg"]), s["respondida_por"].capitalize(),
            analista._QUEM.get(s["ultima_msg_de"], s["ultima_msg_de"]),
            _naive(s["aguardando_desde"]) if s["aguardando_desde"] else None,
            analista._sim_nao(s["sinal_saudacao"]), analista._sim_nao(s["sinal_pediu_nome"]),
            analista._sim_nao(s["sinal_cta"]), analista._sim_nao(s["sinal_followup"]),
            s["media_caracteres_equipe"], s["maior_msg_equipe"],
        ])
    _tabela(ws, analista.COLUNAS_LEADS, linhas,
            [16, 22, 17, 17, 10, 10, 11, 9, 9, 9, 14, 14, 16, 16, 14, 13, 12, 17, 11, 11, 13, 11, 13, 12],
            formatos={3: "dd/mm/yyyy hh:mm", 4: "dd/mm/yyyy hh:mm", 18: "dd/mm/yyyy hh:mm"},
            centralizar=tuple(range(5, 25)))
    ws.row_dimensions[1].height = 48


def _aba_conversas(wb, mensagens):
    ws = wb.create_sheet("Conversas")
    linhas = []
    for m in mensagens:
        if m["tipo"] in ("ferramenta", "sistema") or m["origem"] == "sistema":
            continue
        local = _naive(m["ts"])
        linhas.append([m["telefone"], m["nome"] or "Sem nome", local, local,
                       analista._QUEM.get(m["origem"], m["origem"]), m["tipo"], m["text"]])
    _tabela(ws, analista.COLUNAS_CONVERSAS, linhas, [16, 22, 12, 10, 11, 9, 90],
            formatos={3: "dd/mm/yyyy", 4: "hh:mm:ss"}, centralizar=(3, 4, 5, 6))


def _aba_legenda(wb, metodo):
    ws = wb.create_sheet("Legenda")
    ws.sheet_view.showGridLines = False
    linhas = [("Situação da conversa", None, True)]
    linhas += [
        ("Sucesso", "Agendou aula experimental, combinou visita, matriculou ou pagou.", VERDE),
        ("Em andamento", "A conversa segue viva, sem decisão ainda.", AZUL),
        ("Pede atenção", "A equipe deveria agir hoje (cliente comprador sem data marcada, pedido sem retorno...).",
         AMARELO),
        ("Esperando resposta", "O cliente falou por último e ninguém respondeu ainda.", VERMELHO),
        ("Desistiu", "Disse que não vai ou sumiu (o motivo aparece entre parênteses).", VERMELHO),
        ("Não é venda", "Aluno, fornecedor, engano... fica fora das médias.", CINZA),
        ("", None, None),
        ("Notas (0 a 10)", None, True),
        ("8 a 10", "Muito bom.", VERDE), ("6 a 7,9", "Bom, com espaço para melhorar.", AMARELO),
        ("0 a 5,9", "Dá para melhorar bastante: veja a dica na aba Análises.", VERMELHO),
        ("n/a", "O critério não se aplicou àquela conversa (ex.: não houve objeção). Fica fora da média.", None),
        ("", None, None),
        ("Como a nota é calculada", None, True),
        ("Nota final", "Média dos critérios que se aplicam à conversa.", None),
        ("Tempo da 1ª resposta / médio", "Calculados pelo sistema no horário de atendimento: até 5 min = 10, "
                                          "até 15 = 8, até 30 = 6, até 60 = 4, até 180 = 2, acima disso = 0.", None),
        ("Demais critérios", "Avaliados pela IA lendo a conversa, sempre com base em trechos reais.", None),
        ("Intuito de ligação", "A atendente propôs ligar para o cliente quando fazia sentido.", None),
        ("", None, None),
        ("Termômetro do lead (do mais frio ao mais quente)", None, True),
    ]
    linhas += [(rotulo, desc, cor) for rotulo, desc, cor in TERMOMETRO_INFO.values()]
    linhas += [("", None, None), ("Etapas do script (até onde a conversa chegou)", None, True)]
    linhas += [(rotulo, desc, None) for rotulo, desc in ETAPAS_INFO.values()]
    linhas += [("Funil", "Conta de forma acumulada: quem agendou também foi acolhido, teve o objetivo descoberto "
                         "e recebeu um próximo passo.", None)]
    linhas += [("", None, None), ("Aderência ao script", None, True)]
    linhas += [(f"{rotulo} ({faixa})", desc, cor) for rotulo, faixa, cor, desc in FAIXAS_INFO.values()]
    linhas += [
        ("Como é calculada", "Aderência da conversa = passos feitos ÷ passos que tinham oportunidade. Só contam os passos que tiveram chance de acontecer naquela conversa (ex.: se o preço "
                             "nem foi assunto, \"Preço com valor\" fica fora). A aderência do período soma todos os "
                             "passos feitos e divide por todas as oportunidades. Conversa sem nenhuma oportunidade "
                             "fica fora.", None),
        ("Script - dito × ideal", "Momentos reais da conversa: o que o cliente disse, o que a equipe respondeu e "
                                  "como poderia ser seguindo o script (sugestão da IA).", VERDE_CLARO),
        ("", None, None),
        (f"Técnicas ({metodo['titulo']})", None, True),
    ]
    linhas += [(f"{tid} — {t['nome']}", t["como"], None) for tid, t in metodo["tecnicas"].items()]
    for i, (a, b, extra) in enumerate(linhas, start=1):
        ca = ws.cell(row=i, column=1, value=a)
        if extra is True:
            ca.font = Font(name=FONTE, size=12, bold=True, underline="single", color=ROXO)
            continue
        if not a:
            continue
        cb = ws.cell(row=i, column=2, value=b)
        ca.font, cb.font = Font(name=FONTE, size=10, bold=True), Font(name=FONTE, size=10)
        ca.border = cb.border = BORDA
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(wrap_text=True, vertical="top")
        if extra:
            ca.fill = _fill(extra)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 100
    _config_impressao(ws, None)


# ---------------------------------------------------------------------------
# Montagem
# ---------------------------------------------------------------------------
def construir(agg, analises, sessoes, mensagens, inicio, fim, metodo=None):
    """Monta o .xlsx a partir de dados já carregados e devolve os bytes."""
    metodo = metodo or carregar_metodo()
    wb = openpyxl.Workbook()
    _aba_resumo(wb, agg, inicio, fim)
    if analises:
        _aba_analises(wb, analises, metodo)
        _aba_script(wb, analises, metodo)
        _aba_notas(wb, analises)
        _aba_tecnicas(wb, analises, metodo)
    _aba_leads(wb, sessoes)
    _aba_conversas(wb, mensagens)
    _aba_legenda(wb, metodo)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def gerar(inicio, fim, agora=None):
    """A planilha do período, direto do banco."""
    agora = agora or datetime.now(timezone.utc)
    excluir = analista._ignorados()
    agg, analises = agregar(inicio, fim, agora)
    sessoes = sessoes_repo.do_periodo(inicio, fim, excluir=excluir)
    mensagens = mensagens_repo.do_periodo(inicio, fim, excluir=excluir)
    return construir(agg, analises, sessoes, mensagens, inicio, fim)
