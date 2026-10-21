import io
from datetime import datetime
from decimal import Decimal

import openpyxl
import pytz

from app.services import analista_excel as ex
from app.services.analista_ia import carregar_metodo

TZ = pytz.timezone("America/Sao_Paulo")


def ts(dia, hora, minuto=0):
    return TZ.localize(datetime(2026, 10, dia, hora, minuto))


AGG = {
    "periodo": {}, "conversas": 3, "novas": 2, "retomadas": 1, "followups_da_equipe": 0,
    "msgs_cliente": 8, "msgs_equipe": 6, "com_primeira_resposta": 3, "respondidas_no_limite": 2,
    "primeira_resposta_media_util_seg": 300, "pior_espera_util_seg": 900,
    "respondidas_por": {"atendente": 3, "ninguem": 0}, "sem_resposta": [],
    "analisaveis": 3, "analisadas": 3, "leads": 3, "sucesso": 1, "em_andamento": 1, "perdido": 1, "nao_lead": 0,
    "perdidos_por_motivo": {"preco": 1}, "nota_media": 7.0, "nota_venda": 6.5, "propos_ligacao": 1,
    "script": {
        "titulo": "Script da academia", "fonte": "curso", "conversas": 3, "aderencia": 33, "feitas": 3,
        "oportunidades": 9, "seguiram": 0, "em_parte": 1, "normal": 2,
        "passos": [{"id": "S04", "rotulo": "Descobriu objetivo e dor", "feitas": 3, "oportunidades": 3},
                   {"id": "S09", "rotulo": "Convite com dia e hora", "feitas": 0, "oportunidades": 3}],
        "mais_seguido": {"id": "S04", "nome": "Descobrir objetivo e dor", "feitas": 3, "oportunidades": 3},
        "mais_esquecido": {"id": "S09", "nome": "Convite com dia e hora", "feitas": 0, "oportunidades": 3},
        "termometro": [{"perfil": "comprador", "qtd": 1}, {"perfil": "pesquisador", "qtd": 2}],
        "funil": {"acolhidos": 3, "descobriram": 2, "proximo_passo": 1, "agendaram": 1, "com_etapa": 3},
        "exemplo": None,
    },
    "atencao_total": 1, "pedem_atencao": [], "objecoes": [],
    "por_dia": {
        "2026-10-12": {"conversas": 2, "novas": 1, "retomadas": 1, "sucesso": 1, "em_andamento": 0, "perdido": 1,
                       "nota_media": 8.5, "primeira_resposta_media_util_seg": 240},
        "2026-10-14": {"conversas": 1, "novas": 1, "retomadas": 0, "sucesso": 0, "em_andamento": 1, "perdido": 0,
                       "nota_media": 4.0, "primeira_resposta_media_util_seg": 600},
    },
    "em_aberto_anteriores": 0, "anterior": None, "sintese": None,
}


def analise(nome, tel, dia, desfecho, nota, atencao=False, esperando=None, **extra):
    notas = {"cordialidade": 9, "contexto": None, "primeira_resposta": 10.0}
    base = {
        "nome": nome, "telefone": tel, "inicio": ts(dia, 10), "fim": ts(dia, 11), "nova": True,
        "eh_lead": True, "desfecho": desfecho, "motivo_perda": "preco" if desfecho == "perdido" else None,
        "precisa_atencao": atencao, "aguardando_desde": esperando, "genero": "Feminino",
        "interesse": "muay thai", "nota_final": Decimal(nota), "nota_intuito_venda": Decimal("6.0"),
        "nota_ligacao": None, "primeira_resposta_util_seg": 300, "resumo": "Quer treinar",
        "ponto_forte": "Simpática", "pode_melhorar": "Convidar com dia e hora", "proxima_acao": "Ligar",
        "mensagem_sugerida": "Oi!", "temperatura": "pesquisador",
        "dados": {"notas": notas, "tecnicas_aplicadas": ["S04"], "tecnicas_a_reforcar": ["S09", "S09", "S02"],
                  "etapa_script": "descoberta",
                  "momentos": [{"tecnica": "S09", "cliente_disse": "Posso fazer uma aula teste?",
                                "equipe_respondeu": "Pode sim, quando quiser passar aqui",
                                "pelo_script": "Pode sim! Amanhã às 18h ou quinta às 9h?"}]},
    }
    base.update(extra)
    return base


ANALISES = [
    analise("Maria", "5544991234567", 12, "sucesso", "9.0", temperatura="comprador",
            dados={"notas": {"cordialidade": 9, "contexto": None, "primeira_resposta": 10.0},
                   "tecnicas_aplicadas": ["S01", "S04", "S09"], "tecnicas_a_reforcar": [],
                   "etapa_script": "agendamento", "momentos": []}),
    analise("João", "554488881111", 12, "perdido", "8.0"),
    analise("Carla", "5544977772222", 14, "em_andamento", "4.0", atencao=True),
]
SESSOES = [{
    "telefone": "5544991234567", "nome": "Maria", "inicio": ts(12, 10), "fim": ts(12, 11), "duracao_seg": 3600,
    "nova": True, "iniciada_por": "cliente", "msgs_cliente": 3, "msgs_atendente": 2, "msgs_bot": 0,
    "primeira_resposta_util_seg": 300, "primeira_resposta_seg": 300, "tempo_medio_resposta_util_seg": 300,
    "tempo_medio_resposta_seg": 300, "tempo_max_resposta_util_seg": 300, "respondida_por": "atendente",
    "ultima_msg_de": "atendente", "aguardando_desde": None, "sinal_saudacao": True, "sinal_pediu_nome": False,
    "sinal_cta": True, "sinal_followup": False, "media_caracteres_equipe": 40, "maior_msg_equipe": 80,
}]
MENSAGENS = [
    {"telefone": "5544991234567", "nome": "Maria", "ts": ts(12, 10), "origem": "cliente", "tipo": "texto",
     "text": "Oi, quanto custa?"},
    {"telefone": "5544991234567", "nome": "Maria", "ts": ts(12, 10, 5), "origem": "cliente", "tipo": "texto",
     "text": "=1+1 \x07 com caractere proibido"},
    {"telefone": "5544991234567", "nome": "Maria", "ts": ts(12, 10, 6), "origem": "sistema", "tipo": "sistema",
     "text": "log interno"},
]


def abrir(analises=ANALISES, **kw):
    dados = ex.construir(AGG, analises, SESSOES, MENSAGENS, ts(12, 0), ts(17, 18), carregar_metodo())
    return openpyxl.load_workbook(io.BytesIO(dados))


def cor(cel):
    return cel.fill.start_color.rgb[-6:]


def test_abas_na_ordem_certa():
    assert abrir().sheetnames == ["Resumo", "Análises (IA)", "Script - dito × ideal", "Notas por critério",
                                  "Técnicas", "Leads", "Conversas", "Legenda"]


def test_sem_analises_nao_cria_as_abas_da_ia():
    assert abrir(analises=[]).sheetnames == ["Resumo", "Leads", "Conversas", "Legenda"]


def test_resumo_titulo_blocos_e_dia_a_dia():
    ws = abrir()["Resumo"]
    assert ws["A1"].value == "Brooklyn Academia — Análise do atendimento"
    assert ws["A1"].font.bold and ws["A1"].font.size == 18
    assert ws["A4"].value == "Movimento" and ws["A4"].font.underline == "single" and ws["A4"].font.bold
    assert ws["A5"].value == "Conversas com clientes" and ws["B5"].value == 3 and ws["B5"].font.bold
    rotulos = [c.value for c in ws["A"] if c.value]
    assert "Nota média (0 a 10)" in rotulos and "Propôs ligação para o cliente" in rotulos
    # tabela dia a dia: cabeçalho roxo na linha 5, a partir da coluna D
    assert [c.value for c in ws[5][3:12]][:3] == ["Dia", "Conversas", "Novos"]
    assert cor(ws["D5"]) == ex.ROXO and ws["D5"].font.color.rgb[-6:] == "FFFFFF"
    assert ws["D6"].value == "seg 12/10" and ws["E6"].value == 2
    assert len(ws._charts) == 2                       # conversas por dia + nota média
    nota_tue = [c for c in ws["K"] if c.value == 8.5][0]
    assert cor(nota_tue) == ex.VERDE


def test_resumo_pula_fim_de_semana_sem_conversa():
    ws = abrir()["Resumo"]
    dias = [c.value for c in ws["D"][5:] if c.value]
    assert "sáb 17/10" not in dias and "ter 13/10" in dias      # terça sem conversa aparece zerada


def test_analises_tem_estilo_e_cores():
    ws = abrir()["Análises (IA)"]
    assert [c.value for c in ws[1]][:3] == ["Data", "Nome", "Telefone"]
    assert cor(ws["A1"]) == ex.ROXO and ws["A1"].font.bold and ws["A1"].border.left.style == "thin"
    assert ws["B2"].value == "Maria" and ws["B2"].border.bottom.style == "thin"
    assert ws.freeze_panes == "C2" and ws.auto_filter.ref.startswith("A1:")
    situacoes = {r[1].value: (r[6].value, cor(r[6])) for r in ws.iter_rows(min_row=2)}
    assert situacoes["Maria"] == ("Sucesso", ex.VERDE)
    assert situacoes["João"] == ("Desistiu (preço)", ex.VERMELHO)
    assert situacoes["Carla"] == ("Pede atenção", ex.AMARELO)
    notas = {r[1].value: (r[7].value, cor(r[7])) for r in ws.iter_rows(min_row=2)}
    assert notas["Maria"] == (9.0, ex.VERDE) and notas["Carla"] == (4.0, ex.VERMELHO)
    assert ws["A2"].number_format == "dd/mm/yyyy hh:mm" and ws["H2"].number_format == "0.0"
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToWidth == 1


def test_cliente_esperando_resposta_fica_vermelho():
    ws = abrir([analise("Ana", "55449", 12, "em_andamento", "5.0", esperando=ts(12, 10, 30))])["Análises (IA)"]
    assert ws["G2"].value == "Esperando resposta" and cor(ws["G2"]) == ex.VERMELHO


def test_notas_por_criterio_com_na_e_media():
    ws = abrir()["Notas por critério"]
    cab = [c.value for c in ws[1]]
    assert cab[:4] == ["Data", "Nome", "Telefone", "Nota final"] and "Intuito de ligação" in cab
    col_cord = cab.index("Cordialidade inicial") + 1
    col_ctx = cab.index("Retenção de contexto") + 1
    assert ws.cell(row=2, column=col_cord).value == 9 and cor(ws.cell(row=2, column=col_cord)) == ex.VERDE
    assert ws.cell(row=2, column=col_ctx).value == "n/a"                 # critério que não se aplicou
    ultima = ws.max_row
    assert ws.cell(row=ultima, column=1).value == "Média"
    assert str(ws.cell(row=ultima, column=4).value).startswith("=IFERROR(AVERAGE(D2:D")


def test_tecnicas_contam_conversas_e_a_porcentagem_seguida():
    ws = abrir()["Técnicas"]
    assert [c.value for c in ws[1]] == ["ID", "Técnica (Script da academia)", "Como reconhecer",
                                        "Aplicada em (conversas)", "A reforçar em (conversas)", "% seguida"]
    linhas = {r[0].value: r for r in ws.iter_rows(min_row=2) if r[0].value and str(r[0].value).startswith("S")}
    assert len(linhas) == 11
    valores = lambda tid: tuple(c.value for c in linhas[tid][3:6])
    # S09 duplicado na mesma conversa conta 1 vez; na Maria foi feito
    assert valores("S04") == (3, 0, 1.0) and valores("S09") == (1, 2, 1 / 3) and valores("S01") == (1, 0, 1.0)
    assert valores("S02") == (0, 2, 0.0)
    assert valores("S11") == (0, 0, "n/a")                         # nunca teve oportunidade
    assert cor(linhas["S04"][5]) == ex.VERDE and cor(linhas["S09"][5]) == ex.VERMELHO
    assert linhas["S04"][5].number_format == "0%"


def test_conversas_nao_viram_formula_e_perdem_caractere_proibido():
    ws = abrir()["Conversas"]
    textos = [r[6].value for r in ws.iter_rows(min_row=2)]
    assert len(textos) == 2                                           # log interno ficou de fora
    assert textos[1].startswith(" =1+1") and "\x07" not in textos[1]
    assert ws["G2"].data_type == "s" and ws["G3"].data_type == "s"


def test_legenda_explica_as_cores_o_termometro_as_etapas_e_a_aderencia():
    ws = abrir()["Legenda"]
    valores = [c.value for c in ws["A"] if c.value]
    assert "Sucesso" in valores and "Esperando resposta" in valores and "n/a" in valores
    for termo in ("Tranquilo", "Incomodado", "Pesquisador", "Comparador", "Comprador",
                  "Acolhimento", "Descoberta", "Próximo passo", "Agendamento",
                  "Seguiu o script (70% ou mais)", "Em parte (de 40% a 69%)",
                  "Atendimento normal (abaixo de 40%)", "S09 — Convite com dia e hora"):
        assert termo in valores, termo
    comprador = [c for c in ws["A"] if c.value == "Comprador"][0]
    assert cor(comprador) == ex.TERMOMETRO_INFO["comprador"][2]


def test_analises_tem_termometro_etapa_e_aderencia_colorida():
    ws = abrir()["Análises (IA)"]
    cab = [c.value for c in ws[1]]
    assert cab[11:14] == ["Termômetro", "Etapa do script", "Aderência ao script (%)"]
    assert cab[6:8] == ["Situação", "Nota"]                          # as colunas antigas não mudaram de lugar
    linhas = {r[1].value: r for r in ws.iter_rows(min_row=2)}
    maria, joao = linhas["Maria"], linhas["João"]
    assert (maria[11].value, maria[12].value, maria[13].value) == ("Comprador", "Agendamento", 1.0)
    assert cor(maria[11]) == ex.TERMOMETRO_INFO["comprador"][2] and cor(maria[13]) == ex.VERDE
    assert maria[13].number_format == "0%"
    # 1 feita (S04) de 3 oportunidades (S04, S09, S02): 33% = atendimento normal
    assert joao[11].value == "Pesquisador" and joao[12].value == "Descoberta"
    assert abs(joao[13].value - 1 / 3) < 1e-9 and cor(joao[13]) == ex.VERMELHO


def test_aba_script_dito_x_ideal():
    ws = abrir()["Script - dito × ideal"]
    assert [c.value for c in ws[1]] == ["Data", "Nome", "Telefone", "Técnica", "O cliente disse",
                                        "A equipe respondeu", "Pelo script poderia ser"]
    assert cor(ws["A1"]) == ex.ROXO and ws["A1"].font.bold
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 2                                           # João e Carla (a Maria não tem momento)
    assert linhas[0][1:] == ("João", "554488881111", "S09 — Convite com dia e hora", "Posso fazer uma aula teste?",
                             "Pode sim, quando quiser passar aqui", "Pode sim! Amanhã às 18h ou quinta às 9h?")
    assert ws["G2"].alignment.wrap_text and ws["E2"].border.left.style == "thin"
    assert cor(ws["G2"]) == ex.VERDE_CLARO and ws["A2"].number_format == "dd/mm/yyyy hh:mm"
    assert ws.freeze_panes == "C2" and ws.auto_filter.ref == "A1:G3"


def test_aba_script_sem_momentos_avisa():
    sem = [analise("Ana", "55449", 12, "em_andamento", "5.0",
                   dados={"notas": {}, "tecnicas_aplicadas": [], "tecnicas_a_reforcar": [], "momentos": []})]
    ws = abrir(sem)["Script - dito × ideal"]
    assert ws.max_row == 3 and "Nenhum momento" in ws["B3"].value


def test_resumo_tem_o_bloco_do_script_e_o_funil():
    ws = abrir()["Resumo"]
    valores = {r[0].value: r[1].value for r in ws.iter_rows(min_col=1, max_col=2) if r[0].value}
    assert "Script da academia (lido pela IA)" in valores
    assert valores["Aderência ao script (passos feitos ÷ oportunidades)"] == "33%"
    assert valores["Descobriu objetivo e dor"] == "3 de 3" and valores["Convite com dia e hora"] == "0 de 3"
    assert valores["Passo mais esquecido"] == "Convite com dia e hora"
    assert valores["Termômetro: compradores"] == 1 and valores["Termômetro: pesquisadores"] == 2
    assert valores["Funil do script (até onde as conversas chegaram)"] is None
    assert valores["Etapa"] == "Conversas"
    assert valores["Acolhidos"] == 3 and valores["Descobriram o objetivo"] == "2 · 67%"
    assert valores["Agendaram"] == "1 · 33%"
